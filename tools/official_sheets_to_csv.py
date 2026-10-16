#!/usr/bin/env python3
"""Convert the spreadsheets of ``sources/official/<maker>/`` to CSV, one file per sheet read (DESIGN D123).

The spreadsheets are BIFF (``.xls``) and Office Open XML (``.xlsx``), which the standard library cannot
read, so this one development-time tool (it needs ``xlrd`` and ``openpyxl``, which the ledger does not
depend on) converts the sheets the importer reads, every cell as text, and records each CSV's SHA-256 and
the document and sheet it came from in the manifest. ``rl import official`` reads the CSVs and refuses one
that is not what the manifest says. **A PDF** (KEF's, DESIGN D243) is read as the text of a page, ``page N``, or of
every page, ``pages`` (it needs PyMuPDF): one non-empty line to a row, in the order the library reads the page (``pages``
puts the page number in a first cell), which the importer parses; or as its tables, ``tables`` (Arcam's, D251): a row is
the page, the table's number on it and the cells of one table row, so a table's merged cells are the library's reading.

    python3 tools/official_sheets_to_csv.py marantz|anthem|oppo|kef|jvc|arcam|cambridge|hegel|cary|nad|simaudio|rotel
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

#: maker -> {document: [sheet names read]}
SHEETS = {
    "marantz": {
        "marantz-master-ir.xls": ["AVR Commands", "Direct Vol.", "Zone Speaker Direct Vol.", "AM FM Preset Direct",
                                  "XM SIRIUS Preset Direct", "XM SIRIUS Ch. Direct"],
        "marantz-2014-ir-command-sheet.xls": ["AVR Commands", "Direct Vol."],
        "marantz_fy18_av_sr_nr_ir_code_v02-02072018.xls": ["AVR Commands", "Direct Vol."],
    },
    "oppo": {
        "BDP-103_BDP-103D_Remote_Code_v1.2.xls": ["Remote Code 1", "Remote Code 2", "Remote Code 3", "Notes"],
        "BDP-103_BDP-105_Remote_Code_v1.1.xls": ["Remote Code 1", "Remote Code 2", "Remote Code 3", "Notes"],
        "UDP-203_Remote_Code_v1.2.xls": ["Remote Code 1", "Remote Code 2", "Remote Code 3", "Notes"],
    },
    "anthem": {"AVM-MRXx40-IR-hex-20251202185749500.xlsx": ["MRX x10-x40 remote layout", "MRX x10-x40 IR hex codes"]},
    "jvc": {"DILA_Remote_Code_Emulation_Long_Hex.pdf": ["pages"]},
    "arcam": {
        "AVR380-450-750_RS232_SH256E_2.pdf": ["tables"],
        "AVR5-10-20-30-40-11-21-31-41_RS232_SH289E_F.pdf": ["tables"],
        "SA10-SA20_RS232_SH277E_B.pdf": ["tables"],
        "SA30_RS232_SH306E_4.pdf": ["tables"],
        "CDS50_RS232_SH275E_C.pdf": ["tables"],
        "UDP411_RS232_SH265_E_1.pdf": ["tables"],
        "Solo_IR_Remote_Codes.pdf": ["tables"],
        "Solo_Mini_RS232.pdf": ["pages"],
        "A18-A28-A38_amps_rc.pdf": ["lines"],
        "CD17-CD37_cd_rc.pdf": ["lines"],
        "AVR500-600-AV888_RS232_SH217_E_7.pdf": ["tables"],
        "AVR400_RS232_SH235_E_1.2.pdf": ["tables"],
        "AVR390-550-850-AV860-SR250_RS232_SH274E_D.pdf": ["tables"],
        "BDP100-BDP300_RS232_SH231_E_2.pdf": ["tables"],
        "CDS27_RS232_SH272_E_1.pdf": ["tables"],
        "D33_RS232_SH236_E_D.pdf": ["tables"],
        "ST60_RS232_SH309_C.pdf": ["tables"],
        "Solo_Movie_2.1-5.1_RS232_SH194_3.pdf": ["lines"],
        "Solo_bar_IR_codes_V2.0.pdf": ["lines"],
        "Solo_Movie-Movie_2.1-Music_2_RS232_SH284E_A.pdf": ["tables"],
    },
    "infocus": {"InFocus-IN5310-IRCodes.pdf": ["pages"], "IN5312a_IRCodes.pdf": ["pages"], "IN5316HDa_IRCodes.pdf": ["pages"],         # the four whose table is an image (D298) have the text layer and a transcription
                **{name: ["lines", "transcribed"] for name in ("InFocus-IN110-IRCodes.pdf", "InFocus-IN3118HD-IRCodes.pdf", "IN3114_IN3116_IRCodes.pdf", "InFocus-IN3910-IRCodes.pdf")},
                "SP8604_IRCodes.pdf": ["transcribed"],        # no text layer at all: the whole page is pictures (D301)
                **{name: ["lines"] for name in ("IN5550-IR-Codes.pdf", "InFocus-IN3900LB-IRCodes.pdf")},
                **{name: ["lines"] for name in (
        "540-640-IR_codes.pdf",
        "IN2x_IR_Codes.pdf",
        "IN5140_IRCodes.pdf",
        "IN5148HD-IR-Codes.pdf",
        "IN7x_IR_Codes.pdf",
        "InFocus-IN100-IRCodes.pdf",
        "InFocus-IN1100-IRCodes.pdf",
        "InFocus-IN20-IRCodes.pdf",
        "InFocus-IN2100-IRCodes.pdf",
        "InFocus-IN3100-IRCodes.pdf",
        "InFocus-IN32-IRcodes.pdf",
        "InFocus-IN5300-IRCodes.pdf",
        "InFocus-IN5500-IRCodes.pdf",
        "InFocus-LP120-IRCodes.pdf",
        "InFocus-LP240-LP250-IRCodes.pdf",
        "InFocus-LP600-IRCodes.pdf",
        "InFocus-LP650-IRCodes.pdf",
        "InFocus-LP815-IRCodes.pdf",
        "InFocus-LP840-IRCodes.pdf",
        "InFocus-LS110-IRCodes.pdf",
        "InFocus-SP4800-IRCodes.pdf",
        "InFocus-SP4805-IRCodes.pdf",
        "InFocus-SP5000-IRCodes.pdf",
        "InFocus-SP7200-IRCodes.pdf",
        "InFocus-SP7205-IRCodes.pdf",
        "Mondopad_IR_codes.pdf",
        "X2_IRCodes.pdf",
        "X3_IRCodes.pdf",
    )}},
    "lyngdorf": {"Lyngdorf-IR-Codes.pdf": ["lines"]},
    "harman": {"RS232_SDR35_38_SDP55_58_SH289E_E_2Jun21.pdf": ["tables"], "RS232_Protocol_Documentation.pdf": ["tables"]},
    "yamaha": {name: ["grid"] for name in (
        "Extended_IR_Codes_RX-V685RX-A780-RX-A3080CX-A5200_6.pdf", "Standard_IR_Codes_RX-V685RX-A780-RX-A3080CX-A5200_4.pdf",
        "RX-V385_RX-V485_RX-V585_IR_Codes_2.pdf", "WXA-50_WXC-50_Extended_IR_Codes_1.pdf", "WXA-50_WXC-50_Standard_IR_Codes_1.pdf",
        "XDA-QS5400_XDA-AMP5400_IR_Commands.pdf")},
    "classe": {name: ["cells"] for name in (
        "CAP-2100_IR_codes.pdf", "CDP-100_IR_codes_v1-20.pdf", "CDP-102_202_300_500_and_CDT-300_IR_codes_v1-19.pdf", "CP-500_IR_codes.pdf", "CP-700_IR_codes.pdf",
        "SSP-300_600_IR_codes_v1-20.pdf", "CLASSE_Delta_1_2_Amplifier_IR_Commands.pdf", "CLASSE_SSP-800_IR_codes_v1-24.pdf")} |
        {name: ["pages"] for name in ("CLASSE_CP-800_HEX.pdf", "CLASSE_Delta_1_2_Amplifiers_HEX.pdf", "CLASSE_SSP-800_CT-SSP_HEX.pdf", "CLASSE_SSP-25_HEX.pdf")},
    "cyrus": {name: ["cells"] for name in (
        "ACA75PreXremotecodes.pdf",
        "AV5_AV8remotecodes.pdf",
        "AVMasterremotecodes.pdf",
        "CDremotecodes.pdf",
        "Cyrus3iremotecodes.pdf",
        "Cyrus3remotecodes.pdf",
        "Cyrus5_6_7_8remotecodes.pdf",
        "CyrusPreremotecodes.pdf",
        "dad1remotecodes.pdf",
        "DVD7remotecodes.pdf",
        "FM7FM7.5andFMXremotecodes.pdf",
        "originaltunerremotecodes.pdf",
        "Quattroremotecodes.pdf",
    )},
    "denon": {"AVR4311CI_AVR4311_IR-Codes.pdf": ["words"], "Denon_Master_IR_Hex.xls": ['Main Zone ID#1', 'Main Zone ID#2', 'Main Zone ID#3', 'Main Zone ID#4', 'Zone 2 ID#1', 'Zone 2 ID#2', 'Zone 2 ID#3', 'Zone 2 ID#4', 'Zone 3 ID#1', 'Zone 3 ID#2', 'Zone 3 ID#3', 'Zone 3 ID#4', 'Zone 4 ID#1', 'Zone 4 ID #2', 'Zone 4 ID#3', 'Zone 4 ID#4', 'Analog Tuner ID#1', 'Analog Tuner ID#2', 'Analog Tuner ID#3', 'Analog Tuner ID#4', 'HD Radio ID#1', 'Net Audio ID#1', 'Net Audio ID#2', 'Net Audio ID#3', 'Net Audio ID#4', 'iPod ID#1', 'iPod ID#2', 'iPod ID#3', 'iPod ID#4', 'XM ID#1', 'XM ID#2', 'XM ID#3', 'XM ID#4', 'Sirius ID#1', 'RCD-N7', 'CD Player', 'DBP-1611UD', 'BD Players', 'S5-BD', 'ASD-51NW', 'M33 CD Receiver', 'DRA', 'Sharp Device 2', 'Sharp device 4', 'Sharp Device 12', 'Denon K 4.1', 'Denon K 4.3', 'Denon K 4.5', 'Denon K 4.7', 'Denon K 7.6', 'Denon K 7.7', 'Denon K 7.8']},
    "hitachi": {"2003_and_2004_HITACHI_IR.xls": ["Sheet1", "Sheet2"], "2004_IR_Codes-1.xls": ["2004 Codes", "Additional Codes"],
                **{name: ["cells"] for name in ("2007_Discrete_Codes-DW3.pdf", "2009_Discrete_Codes_L32_and_larger_0211.pdf",
                                                "2009_RC_Layouts_and_IR_Codes_L19_L26.pdf", "AVC_56_Discrete_Codes_UT_Monitor_PDP_Pronto_Hex.pdf")}},
    "emotiva": {name: ["pages"] for name in ("UMC_remote_codes.pdf", "remoteCodes_1-25-11.pdf", "umc200_remote_codes.pdf", "usp_remote_codes.pdf", "erc1_hex.pdf", "ERC-3_IR_CODES.pdf")},
    "cambridge": {"AXR85_100_IR_Codes.pdf": ["tables"], "752BD-CXU-CXUHD_remotes_Sept_2024_v2.pdf": ["tables"]},
    "hegel": {name: ["flow"] for name in (
        "Hegel_Rost_IR_Codes.pdf", "Hegel_H190_IR_Codes.pdf", "Hegel_H360_IR_Codes.pdf", "Hegel_H590_IR_Codes.pdf", "Hegel_H90_IR_Codes.pdf",
        "Hegel_HD12_IR_Codes.pdf", "Hegel_HD25_IR_Codes.pdf", "Hegel_HD30_IR_Codes.pdf", "Hegel_Mohican_IR_Codes.pdf",
        "Hegel_H390_H590_IR_codes_v1.pdf", "Hegel_H95_H120_H190_IR_codes_v2.pdf", "Hegel_P20_P30_IR_codes_v2.pdf")},
    "cary": {"DMS-IR-Remote-Codes.pdf": ["tables"]},
    "nad": {**{name: ["records"] for name in (
        "crestron_cd1.ir", "crestron_cd2.ir", "crestron_cd3.ir", "crestron_cd4.ir", "crestron_cd5.ir", "crestron_dvd1.ir", "crestron_dvd2.ir",
        "crestron_receiver1.ir", "crestron_receiver2_1.ir", "crestron_receiver2_2.ir", "crestron_receiver3.ir", "crestron_receiver4.ir",
        "crestron_tape.ir", "crestron_tuner.ir", "crestron_tv1.ir", "crestron_tv2.ir")},
        "crestron_readme.pdf": ["pages"], **{name: ["cells"] for name in (
        "T175_IR_Commands.pdf", "T765_IR_Commands.pdf", "T775_IR_Commands_2008-02.pdf", "T775_IR_Commands.pdf", "T785_IR_Commands.pdf",
        "T187_IR_Commands.pdf", "T777_IR_Commands.pdf", "T787_IR_Commands.pdf")}},
    "simaudio": {"Simaudio-RC-5-IR-codes-rev1.1.pdf": ["pages"]},
    "rotel": {"Rotel_IR_Codes.html": ["tables"], **{name: ["pages"] for name in (
        "P5-X5_HEX.PDF",
        "P5-P5S2-X5-X5S2_HEX.pdf",
        "Q5_HEX.pdf",
        "X3_HEX.PDF",
        "X3-X3S2_HEX.pdf",
        "X430_HEX.pdf",
        "P5-X5-X5S2_HEX.pdf",
        "CD11_HEX.pdf",
        "CD11-CD11MKII_HEX.pdf",
        "CD14_HEX.pdf",
        "CD14andCD14MKII_HEX.pdf",
        "RCD06SE_HEX.pdf",
        "RCD12_HEX.pdf",
        "RCD1520_HEX.pdf",
        "RCD1570_HEX_0.pdf",
        "RCD1572_HEX.pdf",
        "RCD1572-RCD1572MKII_HEX.pdf",
        "DX3-DX5_HEX.pdf",
        "A10MKII_HEX.pdf",
        "A11_HEX.pdf",
        "A11-A11MKII_HEX.pdf",
        "A12-A14_HEX.pdf",
        "A12-A14-A12MKII-A14MKII_HEX.pdf",
        "A10MKIIandA8_HEX.pdf",
        "RA05SE_HEX.pdf",
        "RA12_HEX.pdf",
        "RA1520_HEX.pdf",
        "RA1570_HEX_0.pdf",
        "RA1572_HEX.pdf",
        "RA1572-RA1572MKII_HEX.pdf",
        "RA1592-RA1592MKII_HEX.pdf",
        "S14_RAS-5000_HEX.pdf",
        "RC1550_HEX.pdf",
        "RC1570_HEX.pdf",
        "RC1572_HEX.pdf",
        "RC1572MKII_HEX.pdf",
        "RC1580_HEX.pdf",
        "RC1590_HEX.pdf",
        "RC1590-RC1590MKII_HEX.pdf",
        "RDD1580_HEX.pdf",
        "RCX1500_HEX.pdf",
        "RAP1580_HEX.pdf",
        "RSP1570_HEX.pdf",
        "RSP1572_HEX.pdf",
        "RSP1576_HEX.pdf",
        "RSP1582_HEX.pdf",
        "RSX1560_HEX.pdf",
        "RSX1562_HEX.pdf",
        "RDG1520_HEX.pdf",
        "RT11_HEX.pdf",
        "RT12_HEX.pdf",
        "T14_HEX.pdf",
    )}},
    "kef": {
        "LS50Wireless-IRcode-INT.pdf": ["page 1", "page 2"],
        "LS50WII_IR-Codes-for-Universal-Remote-Control.pdf": ["page 1"],
        "LS60W_IR_Codes_for_Universal_Remote_Control.pdf": ["page 1"],
        "LSX-IRcode-INT.pdf": ["page 1"],
        "LSX-II_IR-Codes-for-Universal-Remote-Control.pdf": ["page 1"],
        "XIO_IR-Command-Codes-for-Universal-Remote-Control.pdf": ["page 1"],
    },
}


def text(value) -> str:
    """A cell as text: an integral number without ``.0``, anything else as written, trimmed."""
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


class _Table(HTMLParser):
    """The first ``<table>`` of an HTML page (``every``: all of them, each row led by its table's number) as rows of cell text:
    ``<br>`` is `` / ``, ``<sup>x</sup>`` is ``^x``, runs of space one, and with ``links`` an anchor's address follows its text in
    braces."""

    def __init__(self, every: bool = False, links: bool = False) -> None:
        super().__init__()
        self.every, self.links = every, links
        self.rows: list[list[str]] = []
        self.cell: list[str] | None = None
        self.href = ""
        self.depth = 0
        self.tables = 0
        self.done = False

    def handle_starttag(self, tag, attrs):
        if self.done:
            return
        if tag == "table":
            if self.depth == 0:
                self.tables += 1
            self.depth += 1
        elif self.depth and tag == "tr":
            self.rows.append([str(self.tables)] if self.every else [])
        elif self.depth and tag in ("td", "th"):
            self.cell = []
        elif self.cell is not None and tag == "br":
            self.cell.append(" / ")
        elif self.cell is not None and tag == "sup":
            self.cell.append("^")
        elif self.cell is not None and tag == "a" and self.links:
            self.href = dict(attrs).get("href") or ""

    def handle_endtag(self, tag):
        if self.done:
            return
        if tag == "a" and self.cell is not None and self.href:
            self.cell.append(" {" + self.href + "}")
            self.href = ""
        elif tag in ("td", "th") and self.cell is not None:
            self.rows[-1].append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "table" and self.depth:
            self.depth -= 1
            self.done = self.depth == 0 and not self.every

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)


def ir_records(data: bytes) -> list[list[str]]:
    """A Crestron IR driver file (NAD's, D287) as its records: each is ``length, tag, length - 2 bytes`` (the length counts itself and the
    tag) and the file ends with one zero byte. A row is the record's number, its tag and its body, both in hexadecimal."""
    out: list[list[str]] = []
    i = 0
    while i < len(data) - 1:
        n = data[i]
        if n < 2 or i + n > len(data):
            raise ValueError(f"a record of length {n} at byte {i} does not fit the file")
        out.append([str(len(out)), f"{data[i + 1]:02x}", data[i + 2:i + n].hex()])
        i += n
    if data[i:] != b"\x00":
        raise ValueError(f"the file does not end with one zero byte but with {data[i:].hex()!r}")
    return out


def rows_of(path: Path, sheet: str) -> list[list[str]]:
    if path.suffix.lower() == ".ir":
        return ir_records(path.read_bytes())
    if path.suffix.lower() == ".html":
        parser = _Table(every=sheet == "tables", links=sheet == "tables")
        parser.feed(path.read_text(encoding="utf-8"))
        return [[re.sub(r"\s*/\s*$", "", cell).strip() for cell in row] for row in parser.rows]
    if path.suffix.lower() == ".pdf":
        import pymupdf
        pdf = pymupdf.open(str(path))
        if sheet == "tables":       # every table the library finds on every page: page, table number and the cells as text
            return [[str(n), str(i)] + [" ".join((cell or "").split()) for cell in row]
                    for n, page in enumerate(pdf, 1) for i, table in enumerate(page.find_tables().tables, 1)
                    for row in table.extract()]
        if sheet == "flow":         # every page in reading order (D259): a line of text is ``page, text, y, line``, a table row
            out = []                # ``page, row, table, y, cells`` (a cell a merged cell covers is ``<merged>``), sorted by ``y``;
            for n, page in enumerate(pdf, 1):          # a cell much taller than a row is ``<block>`` and its lines are lines of text
                tables = page.find_tables().tables
                items, blocks = [], []
                for i, table in enumerate(tables, 1):
                    for row, cells in zip(table.rows, table.extract()):
                        shown = []
                        for box, cell in zip(row.cells, cells):
                            if box is not None and box[3] - box[1] > 40:
                                blocks.append(box)
                                shown.append("<block>")
                            else:
                                shown.append("<merged>" if cell is None else " ".join(cell.split()))
                        items.append((row.bbox[1], 1, [str(n), "row", str(i), str(round(row.bbox[1]))] + shown))
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        words = " ".join("".join(span["text"] for span in line["spans"]).split())
                        x0, y0 = line["bbox"][0], line["bbox"][1]
                        inside = any(t.bbox[0] - 2 <= x0 <= t.bbox[2] + 2 and t.bbox[1] - 2 <= y0 <= t.bbox[3] + 2 for t in tables)
                        in_block = any(b[0] - 2 <= x0 <= b[2] + 2 and b[1] - 2 <= y0 <= b[3] + 2 for b in blocks)
                        if words and (not inside or in_block):
                            items.append((y0, 0, [str(n), "text", str(round(y0)), words]))
                out.extend(row for _, _, row in sorted(items, key=lambda item: (item[0], item[1])))
            return out
        if sheet == "cells":        # a table with no ruling lines (NAD's, D283): the text of every page as rows of cells at a place: a row is
            out = []                # ``page, y, x, text, x, text...``, the cells of one line (``y`` within 2 points) by their left edge ``x``
            for n, page in enumerate(pdf, 1):
                spans = []
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        for span in line["spans"]:
                            cell = " ".join(span["text"].split())
                            if cell:
                                spans.append((line["bbox"][1], line["bbox"][0], cell))
                spans.sort()
                rows: list[list] = []
                for y, x, cell in spans:
                    if rows and abs(rows[-1][0] - y) <= 2:
                        rows[-1][1].append((x, cell))
                    else:
                        rows.append([y, [(x, cell)]])
                for y, cells in rows:
                    out.append([str(n), str(round(y))] + [v for x, cell in sorted(cells) for v in (str(round(x)), cell)])
            return out
        if sheet == "lines":        # tables with no ruling lines, a cell for each LINE of text (Arcam's, D290): like ``cells``, but the spans of one
            out = []                # line are one cell in the order they are written, so a name set in two fonts stays one name; a line set
            for n, page in enumerate(pdf, 1):       # wholly in italics is written ``<i>text</i>`` (Arcam's CD sheet says what its italics mean)
                lines = []
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        cell = " ".join(" ".join(span["text"] for span in line["spans"]).split())
                        if cell and all(span["flags"] & 2 or "Italic" in span["font"] for span in line["spans"] if span["text"].strip()):
                            cell = f"<i>{cell}</i>"
                        if cell:
                            box = pymupdf.Rect(line["bbox"]) * page.rotation_matrix       # where it is on the page as it is shown (a rotated page: InFocus's spec)
                            lines.append((box.y0, box.x0, cell))
                lines.sort()
                rows: list[list] = []
                for y, x, cell in lines:
                    if rows and abs(rows[-1][0] - y) <= 2:
                        rows[-1][1].append((x, cell))
                    else:
                        rows.append([y, [(x, cell)]])
                for y, cells in rows:
                    out.append([str(n), str(round(y))] + [v for x, cell in sorted(cells) for v in (str(round(x)), cell)])
            return out
        if sheet == "words":        # a table of single characters (Denon's bit matrices, D314): every word of every page by its place, a row is
            out = []                # ``page, y, x, word, x, word...``, the words of one line (``y`` within 3.5 points of the line's first word) by ``x``;
                                    # after a page's lines come its ruled tables' rows, ``page, R, y0, y1``, which say which lines a row's cells hold
            for n, page in enumerate(pdf, 1):
                words = sorted(page.get_text("words"), key=lambda w: (w[1], w[0]))
                rows: list[list] = []
                for w in words:
                    if rows and abs(rows[-1][0] - w[1]) <= 3.5:
                        rows[-1][1].append((w[0], w[4]))
                    else:
                        rows.append([w[1], [(w[0], w[4])]])
                for y, cells in rows:
                    out.append([str(n), str(round(y))] + [v for x, cell in sorted(cells) for v in (str(round(x)), cell)])
                for table in page.find_tables().tables:
                    out.extend([str(n), "R", str(round(row.bbox[1])), str(round(row.bbox[3]))] for row in table.rows)
            return out
        if sheet == "grid":         # a ruled table with merged cells (Yamaha's, D332): each row of each page's tables is ``page, y, one cell for each column``, then the
            out = []                # two columns of names again as the words on the row's own height. A name, a device and a zone are the cell the row is in, and a
            for n, page in enumerate(pdf, 1):       # cell a merged cell covers holds the merged cell's text (a name written once for the rows of its zones, the device and
                words = page.get_text("words")      # the zone of a function); a cell the library finds no merge for is empty, as the document shows it. The codes and the
                for table in page.find_tables().tables:     # models' YES and NO are the words on the row's own height, since a merged cell of two rows may hold two codes,
                    boxes = [c for row in table.rows for c in row.cells if c]       # one on each row's line; and the names are those words too, which say where a merged
                    edges: list[float] = []                 # cell holds the names of several functions (Yamaha's MUTE: ``ON (FULL)``, ``ON (-40dB)``, ``ON (-20dB)``)
                    for x in sorted({b[0] for b in boxes} | {b[2] for b in boxes}):
                        if not edges or x - edges[-1] > 2:
                            edges.append(x)
                    slots = list(zip(edges[:-1], edges[1:]))
                    for row in table.rows:
                        own = [c for c in row.cells if c]
                        if not own:
                            continue
                        small = min(own, key=lambda c: c[3] - c[1])          # the row's own height: its tallest cells may span several rows
                        middle = (small[1] + small[3]) / 2

                        def text_of(box) -> str:
                            inside = sorted((w for w in words if box and box[0] - 1 <= (w[0] + w[2]) / 2 <= box[2] + 1 and box[1] - 1 <= (w[1] + w[3]) / 2 <= box[3] + 1),
                                            key=lambda w: (round(w[1] / 3), w[0]))
                            return " ".join(w[4] for w in inside)

                        cells, banded = [], []
                        for column, (a, b) in enumerate(slots):
                            band = (a, small[1], b, small[3])
                            if column in (0, 1, 2, 3, 4):
                                box = next((c for c in boxes if c[0] - 1 <= (a + b) / 2 <= c[2] + 1 and c[1] - 1 <= middle <= c[3] + 1), None)
                            else:
                                box = band
                            cells.append(text_of(box))
                            if column in (1, 2):
                                banded.append(text_of(band))
                        out.append([str(n), str(round(small[1]))] + cells + banded)
            return out
        if sheet == "pages":        # every page: a row is the page number and one line of its text
            return [[str(n), line.strip()] for n, page in enumerate(pdf, 1) for line in page.get_text("text").splitlines()
                    if line.strip()]
        return [[line.strip()] for line in pdf[int(sheet.removeprefix("page ")) - 1].get_text("text").splitlines() if line.strip()]
    if path.suffix.lower() == ".xls":
        import xlrd
        ws = xlrd.open_workbook(str(path)).sheet_by_name(sheet)
        return [[text(c.value) for c in ws.row(r)] for r in range(ws.nrows)]
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    return [[text(c) for c in row] for row in ws.iter_rows(values_only=True)]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("maker", choices=sorted(SHEETS))
    parser.add_argument("--root", default="sources/official")
    args = parser.parse_args(argv)
    target = Path(args.root) / args.maker
    manifest_path = target / "MANIFEST.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["sheets"] = {}
    (target / "csv").mkdir(exist_ok=True)
    for document, sheets in SHEETS[args.maker].items():
        for sheet in sheets:
            name = re.sub(r"[^A-Za-z0-9._-]+", "_", f"{Path(document).stem}__{sheet}") + ".csv"
            out = target / "csv" / name
            if sheet == "transcribed":      # a table that is an image, written down by eye (D298): the file is the work and is kept as it is
                if not out.is_file():
                    raise SystemExit(f"{out}: the transcription of {document} is not there; it is written by hand and this tool does not make it")
                with out.open(encoding="utf-8", newline="") as handle:
                    rows = list(csv.reader(handle))
                manifest["sheets"][name] = {"document": document, "sheet": sheet, "sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
                                            "rows": len(rows), "transcribed": "by eye from the table's image in the PDF"}
                print(f"{name}: {len(rows)} rows (transcribed, kept)", file=sys.stderr)
                continue
            rows = rows_of(target / document, sheet)
            while rows and not any(rows[-1]):
                rows.pop()
            with out.open("w", encoding="utf-8", newline="") as handle:
                csv.writer(handle, lineterminator="\n").writerows(rows)
            manifest["sheets"][name] = {"document": document, "sheet": sheet,
                                        "sha256": hashlib.sha256(out.read_bytes()).hexdigest(), "rows": len(rows)}
            print(f"{name}: {len(rows)} rows", file=sys.stderr)
    manifest["sheets"] = dict(sorted(manifest["sheets"].items()))
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
